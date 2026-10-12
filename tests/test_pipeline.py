import json
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from recon81.core import Inventory
from recon81.importers import load_records, native_record

PROJECT = Path(__file__).resolve().parents[1]


class PipelineTests(unittest.TestCase):
    def test_native_roundtrip_keeps_variant_id(self):
        for r in load_records(PROJECT/'examples/demo.har'):
            clean = native_record(r)
            self.assertEqual(r['request_id'], clean['request_id'])
            self.assertEqual(r['observation_id'], clean['observation_id'])

    def test_js_pipeline_retains_post_and_excludes_post_from_baseline_get_list(self):
        if not shutil.which('jq'):
            self.skipTest('jq is required for the legacy Bash pipeline')
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            results = root/'results'; results.mkdir()
            jsdir = results/'example.test_js_files_00'; jsdir.mkdir()
            (jsdir/'abc.js').write_text('/* dummy local JS fixture */')
            (results/'example.test_js_mapping.txt').write_text('abc.js\thttps://example.test/static/main.js\n')
            records = [{'url':'/api/post','method':'POST','bodyParams':['userId','password'],'headers':{'Content-Type':'application/json','Authorization':'NEVER_SAVE'}}, {'url':'/api/get','method':'GET'}, {'url':'/api/unknown'}]
            fake = root/'bin'; fake.mkdir()
            tool = fake/'jsluice'
            tool.write_text('#!/usr/bin/env python3\nprint('+repr('\n'.join(json.dumps(x) for x in records))+')\n')
            tool.chmod(0o700)
            (root/'sanitize_js_requests.py').symlink_to(PROJECT/'sanitize_js_requests.py')
            env = dict(os.environ, PATH=str(fake)+os.pathsep+os.environ['PATH'], PYTHONPATH=str(PROJECT))
            subprocess.run(['bash',str(PROJECT/'scan_linkfinder.sh'),'00'],cwd=root,env=env,check=True,capture_output=True)
            text = (results/'example.test_js_requests_00.jsonl').read_text()
            self.assertNotIn('NEVER_SAVE',text)
            stored = [json.loads(x) for x in text.splitlines()]
            self.assertEqual({'POST','GET','UNKNOWN'},{r['method'] for r in stored})
            legacy = (results/'example.test_linkfinder_00.txt').read_text()
            self.assertNotIn('/api/post',legacy)
            self.assertIn('/api/get',legacy)
            scope=root/'scope.txt'; scope.write_text('example.test\n')
            db=root/'inventory.db'; report=root/'report'
            subprocess.run(['python',str(PROJECT/'collect_request_inventory.py'),'--scope-file',str(scope),'--results',str(results),'--db',str(db),'--output',str(report),'--xlsx'],env=env,check=True,capture_output=True)
            with Inventory(db) as inv:
                self.assertEqual(3,len(list(inv.records())))
            self.assertTrue((report/'index.html').is_file())

    def test_global_mixer_accepts_katana_and_does_not_probe_secret_findings(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); (root/'targets.txt').write_text('example.test\n'); results=root/'results'; results.mkdir()
            (results/'example.test_katana_00.txt').write_text('https://example.test/api/fromkatana\n')
            (results/'example.test_trufflehog_00.txt').write_text('https://example.test/secretFinding\n')
            subprocess.run(['python',str(PROJECT/'global_mixer.py')],cwd=root,check=True,capture_output=True)
            data=''.join(p.read_text() for p in (root/'chunks').glob('*.txt'))
            self.assertIn('/api/fromkatana',data)
            self.assertNotIn('/secretFinding',data)


if __name__=='__main__':
    unittest.main()

class LegacyReportTests(unittest.TestCase):
    def test_legacy_workbook_includes_request_sheets_and_masks_url_secrets(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); (root/'targets.txt').write_text('example.test\n'); results=root/'results'; results.mkdir()
            (results/'example.test_gau_00.txt').write_text('https://example.test/api/a?token=NEVER_SAVE_URL_SECRET&id=1\n')
            (results/'example.test_trufflehog_00.txt').write_text('abc.js\t[TokenDetector] [REDACTED]\n')
            (results/'example.test_js_mapping.txt').write_text('abc.js\thttps://example.test/static/main.js\n')
            env=dict(os.environ,PYTHONPATH=str(PROJECT))
            subprocess.run(['python',str(PROJECT/'request_inventory.py'),'import',str(PROJECT/'examples/demo.har'),'--scope','example.test'],cwd=root,env=env,check=True,capture_output=True)
            script='''import sys, types
# The network-only requests package is replaced for this entirely offline test.
sys.modules['requests'] = types.ModuleType('requests')
sys.modules['aiohttp'] = types.ModuleType('aiohttp')
import txt_to_excel as report
report.run_duckduckgo_dorking = lambda targets: []
async def empty_analysis(items): return {}
report.analyze_all_subdomains = empty_analysis
report.build_advanced_excel_report()
'''
            subprocess.run(['python','-c',script],cwd=root,env=env,check=True,capture_output=True)
            from openpyxl import load_workbook
            wb=load_workbook(next((root/'reports').glob('passive_recon_report_*.xlsx')))
            self.assertIn('Requests v8.1',wb.sheetnames)
            self.assertIn('Structures v8.1',wb.sheetnames)
            self.assertIn('Secrets v8.1',wb.sheetnames)
            self.assertNotIn('Gemini',str(wb.sheetnames))
            all_cells=' '.join(str(c.value) for sheet in wb for row in sheet for c in row)
            self.assertNotIn('NEVER_SAVE_URL_SECRET',all_cells)
            self.assertNotIn('https://example.test/[TokenDetector]',all_cells)
            self.assertIn('TokenDetector',all_cells)
