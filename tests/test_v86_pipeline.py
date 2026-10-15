"""Regression fixtures for full discovery retention and export agreement."""
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from recon81.core import Inventory, Scope, normalize_record
from recon81.network_policy import NetworkPolicy
from recon81.pipeline import endpoint_signature, execution_plan, ingest_saved, normalize_dynamic_path

PROJECT=Path(__file__).resolve().parents[1]


class CanonicalPipelineTests(unittest.TestCase):
    def test_transport_ledger_is_not_imported_as_target_response(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            results = root / 'results'
            results.mkdir()
            (results / 'httpx_results_00.json').write_text(json.dumps({'url': 'https://example.test/api/a', 'status_code': 200}) + '\n')
            (results / 'httpx_results_00.network.json').write_text(json.dumps({'guard': {'requests': 1}, 'events': [{'status': 502}]}, indent=2))
            with Inventory(root / 'inventory.db') as inventory:
                stats = ingest_saved(results, Scope(['example.test']), inventory)
                records = list(inventory.records())
            self.assertEqual(0, stats['parse_errors'])
            self.assertEqual([200], [record['response']['status'] for record in records])

    def run_script(self, root, name, *args):
        result=subprocess.run([sys.executable,str(PROJECT/name),*map(str,args)],cwd=root,env=dict(os.environ,PYTHONPATH=str(PROJECT)),capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr)
        return result

    def test_distinct_siblings_survive_mixer_history_and_every_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); results=root/'results'; results.mkdir()
            (root/'targets.txt').write_text('example.test\n')
            urls={'https://example.test/api/'+name for name in ('alpha','beta','gamma','delta','epsilon','zeta','eta','theta')}
            archive=results/'example.test_gau_00.txt'
            archive.write_text('\n'.join(sorted(urls))+'\n')
            self.run_script(root,'global_mixer.py')
            queue={line for path in (root/'chunks').glob('*.txt') for line in path.read_text().splitlines()}
            self.assertEqual(queue,urls)
            self.assertEqual(len(list((root/'chunks').glob('chunk_*.txt'))),20)
            first_assignment={p.name:p.read_text() for p in (root/'chunks').glob('*.txt')}
            self.run_script(root,'global_mixer.py')
            self.assertEqual(first_assignment,{p.name:p.read_text() for p in (root/'chunks').glob('*.txt')})
            self.run_script(root,'txt_to_excel.py')
            from openpyxl import load_workbook
            wb=load_workbook(next((root/'reports').glob('passive_recon_report_*.xlsx')))
            self.assertEqual(urls,{row[3] for row in wb['example.test'].iter_rows(min_row=3,values_only=True)})
            self.assertEqual(urls,{row[2] for row in wb['Requests v8.1'].iter_rows(min_row=2,values_only=True)})
            wb.close()
            exported=[json.loads(line) for line in (root/'reports/requests/requests.jsonl').read_text().splitlines()]
            self.assertEqual(urls,{r['url'] for r in exported})
            postman=json.loads(next((root/'reports').glob('postman_collection_*.json')).read_text())
            native=json.loads((root/'reports/requests/postman_review.json').read_text())
            self.assertEqual(postman,native)
            self.assertEqual(urls,{i['request']['url'] for i in postman['item']})
            # Removal from today's collector input does not delete historical data.
            archive.unlink()
            (root/'reports/new_subdomains.txt').write_text('stale.example.test\n')
            self.run_script(root,'txt_to_excel.py')
            self.assertEqual((root/'reports/new_subdomains.txt').read_text(),'')
            with sqlite3.connect(root/'reports/recon_history.db') as conn:
                self.assertEqual(urls,{row[0] for row in conn.execute('SELECT url FROM master_urls')})
                self.assertEqual(8,conn.execute('SELECT passive_tot FROM target_stats WHERE target=?',('example.test',)).fetchone()[0])
            with Inventory(root/'reports/request_inventory.sqlite3') as inv:
                self.assertEqual(urls,{r['url'] for r in inv.records()})
            self.assertEqual(urls,set((root/'reports/master_url_db.txt').read_text().splitlines()))

    def test_sampling_is_only_an_explicit_execution_view(self):
        scope=Scope(['example.test'])
        records=[normalize_record('https://example.test/api/items/'+str(i),'UNKNOWN',source='gau',observed=False) for i in range(100,111)]
        records.extend([normalize_record('https://example.test/api/delete1','UNKNOWN',source='gau',observed=False),normalize_record('https://example.test/api/create','POST',observed=False)])
        selected,deferred,summary=execution_plan(records,NetworkPolicy(scope))
        self.assertEqual(len(records),13)
        self.assertEqual(len(selected),5)
        self.assertEqual(summary['deferred'],8)
        self.assertEqual(summary['deferred_by_reason']['signature_sample_limit'],6)
        self.assertIn('method_not_baseline_get',{r['reason'] for r in deferred})
        self.assertEqual(execution_plan(list(reversed(records)),NetworkPolicy(scope)),(selected,deferred,summary))
        self.assertEqual(set(r['url'] for r in records),set(selected)|{r['url'] for r in deferred})

    def test_signature_keeps_literal_route_method_and_query_multiplicity(self):
        self.assertEqual(normalize_dynamic_path('/administration/configuration'),'/administration/configuration')
        self.assertNotEqual(endpoint_signature('https://example.test/api/alpha'),endpoint_signature('https://example.test/api/beta'))
        self.assertNotEqual(endpoint_signature('https://example.test/api/alpha','GET'),endpoint_signature('https://example.test/api/alpha','POST'))
        self.assertNotEqual(endpoint_signature('https://example.test/api/alpha?id=1'),endpoint_signature('https://example.test/api/alpha?id=1&id=2'))
        self.assertEqual(endpoint_signature('https://example.test/api/100'),endpoint_signature('https://example.test/api/101'))

    def test_bad_lines_are_counted_and_later_evidence_survives(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); results=root/'results'; results.mkdir()
            record=normalize_record('https://example.test/api/a','POST',observed=False)
            (results/'example.test_js_requests_00.jsonl').write_text('{bad secret=SYNTHETIC_PRIVATE}\n'+json.dumps(record)+'\n')
            (results/'httpx_results_00.json').write_text('{broken}\n'+json.dumps({'url':'https://example.test/api/b','status_code':403})+'\n')
            with Inventory(root/'inventory.sqlite3') as inv:
                stats=ingest_saved(results,Scope(['example.test']),inv)
                self.assertTrue(stats['partial']); self.assertEqual(stats['parse_errors'],2)
                self.assertEqual(len(list(inv.records())),2)
                self.assertNotIn('SYNTHETIC_PRIVATE',json.dumps(stats))
                again=ingest_saved(results,Scope(['example.test']),inv)
                self.assertEqual(again['imported'],0)
                self.assertEqual(again['unchanged'],2)

    def test_capture_receipts_keep_repeated_events_without_reimport_inflation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); results=root/'results'; results.mkdir()
            scope=Scope(['example.test'])
            rows=[{'url':'https://example.test/api/a','status_code':200,'timestamp':f'2026-10-15T00:00:0{i}+00:00'} for i in (1,2,3)]
            source=results/'httpx_results_00.json'
            source.write_text('\n'.join(json.dumps(r) for r in rows)+'\n')
            with Inventory(root/'inventory.sqlite3') as inv:
                ingest_saved(results,scope,inv)
                self.assertEqual(len(list(inv.records())),1)
                events=list(inv.capture_events())
                self.assertEqual(len(events),3)
                self.assertEqual([e['timestamp'] for e in events],[f'2026-10-15T00:00:0{i}.000000+00:00' for i in (1,2,3)])
                ingest_saved(results,scope,inv)
                self.assertEqual(len(list(inv.capture_events())),3)
                source.write_text(source.read_text()+json.dumps(dict(rows[0],timestamp='2026-10-15T00:00:04+00:00'))+'\n')
                ingest_saved(results,scope,inv)
                self.assertEqual(len(list(inv.capture_events())),4)

    def test_wildcard_archive_and_rich_js_records_share_inventory(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); results=root/'results'; results.mkdir()
            (results/'wild_example.test_gau_00.txt').write_text('https://api.example.test/a\nhttps://elsewhere.test/b\n')
            record=normalize_record('https://api.example.test/c','GET',observed=False,source='jsluice')
            (results/'wild_example.test_js_requests_00.jsonl').write_text(json.dumps(record)+'\n')
            (results/'wild_example.test_linkfinder_00.txt').write_text(record['url']+'\n')
            with Inventory(root/'inventory.sqlite3') as inv:
                stats=ingest_saved(results,Scope(['*.example.test']),inv)
                self.assertEqual(stats['out_of_scope'],1)
                records=list(inv.records())
                self.assertEqual(len(records),2)
                self.assertEqual({r['method'] for r in records},{'UNKNOWN','GET'})


if __name__=='__main__': unittest.main()
