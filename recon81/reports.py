"""Standalone local reports. No remote scripts, fonts or telemetry."""
import json
import hashlib
import html
import os
import re
import shutil
import tempfile
import uuid
from collections import Counter
from pathlib import Path
from .core import private_write


def review_candidates(record):
    text = ' '.join(record['body']['fields']).lower()
    flags = []
    if any(x in text for x in ('userid', 'user_id', 'owner', 'accountid', 'role', 'permission')):
        flags.append('인가 / 객체 소유권 검토')
    if any(x in text for x in ('url', 'redirect', 'callback', 'destination', 'path', 'filename')):
        flags.append('URL / 파일 입력 처리 검토')
    if record['body']['kind'] == 'multipart':
        flags.append('파일 업로드 처리 검토')
    if record['method'] == 'UNKNOWN':
        flags.append('메소드 확인 필요')
    if not record['observed']:
        flags.append('실제 요청 확보 필요')
    return flags


def postman(records):
    items = []
    seen = set()
    for r in records:
        context=r.get('context') or {}
        key=(r['request_id'],context.get('actor',''),context.get('tenant',''))
        if key in seen:
            continue
        seen.add(key)
        request = {'method': r['method'], 'header': r['headers'], 'url': r['url'], 'description': '검토 전용. 민감값은 마스킹되어 있으며 자동 실행이 차단되어 있습니다. 계정: '+context.get('actor','미지정')+' / 조직: '+context.get('tenant','미지정')+' / 관찰된 요청: ' + str(r['observed'])}
        body = r['body']
        if body['kind'] in ('json', 'form', 'xml', 'text'):
            request['body'] = {'mode': 'raw', 'raw': body['text']}
            if body['kind'] == 'json':
                request['body']['options'] = {'raw': {'language': 'json'}}
        items.append({'name': r['method'] + ' ' + r['url'], 'request': request})
    return {'info': {'name': 'Passive Scan v8.2 — 요청 검토 전용', 'schema': 'https://schema.getpostman.com/json/collection/v2.1.0/collection.json', 'description': '자동 재전송 금지. 사용 전 scope·세션·body를 직접 확인하세요.'}, 'event': [{'listen': 'prerequest', 'script': {'type': 'text/javascript', 'exec': ['throw new Error("검토 전용 컬렉션입니다. 요청을 보내려면 복사본에서 이 차단 스크립트를 직접 제거하세요.");']}}], 'item': items}


from .dashboard import HTML


def add_excel_sheets(workbook, records):
    from openpyxl.styles import Alignment, Font, PatternFill
    for name, summary in [('Requests v8.1', False), ('Structures v8.1', True)]:
        if name in workbook.sheetnames:
            del workbook[name]
        ws = workbook.create_sheet(name)
        ws.append(['No', 'Method', 'URL', '관찰 상태', 'Body 필드', 'Body 값 (마스킹)', '응답 코드', '출처', '구조 ID', '요청 ID', '검토 후보'])
        seen = set()
        for r in records:
            if summary and r['structure_id'] in seen:
                continue
            seen.add(r['structure_id'])
            def cell(value):
                value = str(value)
                if len(value) > 32000:
                    value = value[:31900] + '\n[Excel 표시 제한: 전체 body는 JSONL/HTML에서 확인]'
                if value.startswith(('=', '+', '-', '@')):
                    value = "'" + value
                # XML 1.0 cell content cannot contain these controls.
                return ''.join(x for x in value if ord(x) >= 32 or x in '\t\n\r')
            ws.append([ws.max_row, r['method'], cell(r['url']), '실제 요청' if r['observed'] else '정적 추출', cell(', '.join(r['body']['fields'])), cell(r['body']['text']), (r.get('response') or {}).get('status', ''), cell(r['source']), r['structure_id'], r['request_id'], cell(', '.join(review_candidates(r)))])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.fill = PatternFill('solid', fgColor='193B5B')
            c.font = Font(color='FFFFFF', bold=True)
        widths = [8, 14, 65, 16, 45, 65, 12, 20, 24, 24, 35]
        for i, width in enumerate(widths, 1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(i)].width = width
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical='top', wrap_text=True)


def add_analysis_sheets(workbook,records,analysis):
    from openpyxl.styles import Font,PatternFill,Alignment
    def cell(value):
        if isinstance(value,(dict,list)): value=json.dumps(value,ensure_ascii=False)
        value=str(value)[:32000]
        value=''.join(c for c in value if ord(c)>=32 or c in '\n\t\r')
        return "'"+value if value.startswith(('=','+','-','@')) else value
    tables={
        'Candidates v8.2':(['Observation ID','Method','URL','Actor','Tenant','Scenario','Priority','Score','Review state','Reasons','Missing','Review note'],[[r['observation_id'],r['method'],r['url'],r['context'].get('actor',''),r['context'].get('tenant',''),r['context'].get('scenario',''),r['triage']['priority'],r['triage']['score'],r['review']['state'],r['triage']['hints'],r['triage']['missing'],r['review'].get('note','')] for r in records]),
        'Coverage v8.2':(['Method','Route','Records','Observed','Actors','States','Missing'],[[r['method'],r['route'],r['records'],r['observed'],r['actors'],r['states'],r['missing']] for r in analysis['coverage']]),
        'Compare v8.2':(['Method','URL','Actor A','Actor B','Evidence A','Evidence B','Comparable','Differences','Reason'],[[p['method'],p['url'],p['left_actor'],p['right_actor'],p['left_id'],p['right_id'],p['comparable'],p['differences'],p['reason']] for p in analysis['comparisons']]),
        'Flows v8.2':(['Actor','Tenant','Scenario','Capture','Steps','Inferred links'],[[f['actor'],f['tenant'],f['scenario'],f['capture_id'],f['steps'],f['links']] for f in analysis['flows']])}
    for name,(header,rows) in tables.items():
        if name in workbook.sheetnames: del workbook[name]
        ws=workbook.create_sheet(name); ws.append(header)
        for row in rows: ws.append([cell(x) for x in row])
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for c in ws[1]: c.font=Font(color='FFFFFF'); c.fill=PatternFill('solid',fgColor='193B5B')
        for column in ws.columns: ws.column_dimensions[column[0].column_letter].width=32
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)


def _file_digest(path):
    digest=hashlib.sha256()
    with open(path,'rb') as stream:
        for chunk in iter(lambda:stream.read(65536),b''): digest.update(chunk)
    return digest.hexdigest()


def _private_chunks(path, chunks):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True,mode=0o700)
    fd, temporary=tempfile.mkstemp(dir=path.parent)
    try:
        with os.fdopen(fd,'w',encoding='utf-8') as stream:
            for chunk in chunks: stream.write(chunk)
        os.replace(temporary,path); path.chmod(0o600)
    finally:
        if os.path.exists(temporary): os.unlink(temporary)


def _private_json(path, value):
    _private_chunks(path,json.JSONEncoder(ensure_ascii=False,indent=2).iterencode(value))


def _private_jsonl(path, records):
    def chunks():
        encoder=json.JSONEncoder(ensure_ascii=False)
        for record in records:
            yield from encoder.iterencode(record)
            yield '\n'
    _private_chunks(path,chunks())


def _oversized_evidence_page(records):
    links=''.join('<li><a href="../evidence/'+r['observation_id']+'.json">'+html.escape(r['method']+' '+r['url'][:512])+'</a></li>' for r in records)
    return """<!doctype html><html lang="ko"><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="default-src 'none'; base-uri 'none'; form-action 'none'"><title>큰 증거 파일</title><h1>큰 증거 파일</h1><p>이 관찰은 HTML 표시 예산을 초과합니다. 전체 증거를 생략 없이 JSON으로 열 수 있습니다. 전체 비교는 analysis.json에 있습니다.</p><a href="../index.html">전체 보고서</a><ul>"""+links+'</ul></html>'


def _generate(records, output, xlsx=False, ignore_fields=(), capture_events=None, resource_headers=()):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True, mode=0o700)
    from .analysis import analyze
    from .core import Scope
    from .network_policy import NetworkPolicy
    from urllib.parse import urlsplit
    from review_queue import build
    enriched,analysis=analyze(records,ignore_fields,capture_events=capture_events,resource_headers=resource_headers)
    enriched=[dict(r,review_candidates=review_candidates(r)) for r in enriched]
    _private_jsonl(output / 'requests.jsonl', enriched)
    _private_json(output / 'postman_review.json', postman(records))
    _private_json(output/'analysis.json', analysis)
    hosts={urlsplit(r['url']).hostname for r in records if urlsplit(r['url']).hostname}
    queue=build(records,NetworkPolicy(Scope(hosts))) if hosts else {'schema_version':'8.2','items':[]}
    _private_json(output/'review_queue.json',queue)
    _write_dashboards(output, enriched, analysis, ignore_fields, resource_headers, capture_events)
    for r in enriched:
        private_write(output / 'evidence' / (r['observation_id'] + '.json'), json.dumps(r, ensure_ascii=False, indent=2))
    if xlsx:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = 'Summary'
        ws.append(['항목', '개수'])
        ws.append(['전체 기록', len(records)])
        ws.append(['실제 요청', sum(r['observed'] for r in records)])
        ws.append(['정적 추출', sum(not r['observed'] for r in records)])
        for method, count in sorted(Counter(r['method'] for r in records).items()):
            ws.append([method, count])
        ws.column_dimensions['A'].width = 28
        ws.column_dimensions['B'].width = 16
        add_excel_sheets(wb, records)
        add_analysis_sheets(wb,enriched,analysis)
        target = output / 'request_inventory.xlsx'
        # Reserve privately before openpyxl replaces the content.
        private_write(target, '')
        wb.save(target)
    if capture_events is not None:
        _private_jsonl(output / 'capture_events.jsonl', capture_events)
    files = {path.relative_to(output).as_posix(): _file_digest(path)
             for path in sorted(output.rglob('*')) if path.is_file()}
    private_write(output / 'manifest.json', json.dumps({
        'schema_version': '8.6', 'owner': 'passive-scan-report', 'run_id': uuid.uuid4().hex, 'files': files}, indent=2))
    return len(records)


DASHBOARD_MAX_RECORDS = 1000
DASHBOARD_MAX_BYTES = 8 * 1024 * 1024


def _dashboard_json(records, analysis):
    return json.dumps({'records': records, 'analysis': analysis}, ensure_ascii=False).replace('<', '\\u003c').replace('&', '\\u0026')


def _write_dashboards(output, records, analysis, ignore_fields, resource_headers, capture_events=None):
    """Large reports retain complete JSON evidence and bounded static HTML pages."""
    record_sizes = {r['observation_id']:len(json.dumps(r,ensure_ascii=False).encode()) for r in records}
    budget = max(1024, DASHBOARD_MAX_BYTES-len(HTML.encode())-2048)
    if len(records) <= DASHBOARD_MAX_RECORDS and sum(record_sizes.values()) <= budget:
        data = _dashboard_json(records, analysis)
        if len(data.encode()) <= budget:
            private_write(output / 'index.html', HTML.replace('__DATA__', data))
            return
    from .analysis import analyze
    events_by_id={}
    for event in capture_events or ():
        events_by_id.setdefault(event['observation_id'],[]).append(event)
    def page_analysis(selected):
        events=None if capture_events is None else [event for record in selected for event in events_by_id.get(record['observation_id'],[])]
        return analyze(selected,ignore_fields,capture_events=events,resource_headers=resource_headers)
    chunks = []; chunk = []; size = 0
    # Leave space for derived comparison/coverage metadata in each page.
    for record in records:
        record_size = record_sizes[record['observation_id']]
        if chunk and (len(chunk) >= DASHBOARD_MAX_RECORDS or size + record_size > DASHBOARD_MAX_BYTES // 2):
            chunks.append(chunk); chunk = []; size = 0
        chunk.append(record); size += record_size
    if chunk: chunks.append(chunk)
    links = []
    pending=list(chunks); index=0
    while pending:
        chunk=pending.pop(0)
        subset, details = page_analysis(chunk)
        data=_dashboard_json(subset, details)
        if len(data.encode())>budget and len(chunk)>1:
            middle=len(chunk)//2; pending[0:0]=[chunk[:middle],chunk[middle:]]; continue
        index+=1; filename = 'pages/requests-' + str(index) + '.html'
        if len(data.encode())>budget:
            content=_oversized_evidence_page(chunk)
        else:
            content = HTML.replace('__DATA__', data)
            content = content.replace('<main>', '<main><p><a href="../index.html">전체 보고서</a> · 분할 페이지의 비교·흐름은 이 페이지의 관찰 범위입니다. 전체 분석은 analysis.json에 보존됩니다.</p>', 1)
        private_write(output / filename, content)
        links.append((filename, '요청 ' + str(index) + ' · ' + str(len(chunk)) + '개 기록'))
    # Preserve cross-page account comparisons as bounded, self-contained views.
    by_id = {r['observation_id']: r for r in records}
    pair_chunks = []; pair_chunk = []; pair_ids = set(); pair_size = 0
    for pair in analysis['comparisons']:
        ids = {pair['left_id'], pair['right_id']}
        additional = sum(record_sizes[key] for key in ids - pair_ids)
        if pair_chunk and (len(pair_chunk) >= 30 or pair_size + additional > DASHBOARD_MAX_BYTES // 2):
            pair_chunks.append(pair_chunk); pair_chunk = []; pair_ids = set(); pair_size = 0
            additional = sum(record_sizes[key] for key in ids)
        pair_chunk.append(pair); pair_ids.update(ids); pair_size += additional
    if pair_chunk: pair_chunks.append(pair_chunk)
    pending=list(pair_chunks); index=0
    while pending:
        pairs=pending.pop(0)
        ids = {p[key] for p in pairs for key in ('left_id', 'right_id')}
        selected=[by_id[key] for key in sorted(ids)]
        subset, details = page_analysis(selected)
        details['comparisons'] = pairs
        data=_dashboard_json(subset, details)
        if len(data.encode())>budget and len(pairs)>1:
            middle=len(pairs)//2; pending[0:0]=[pairs[:middle],pairs[middle:]]; continue
        index+=1; filename = 'pages/comparisons-' + str(index) + '.html'
        content = (_oversized_evidence_page(selected) if len(data.encode())>budget else HTML.replace('__DATA__',data).replace('<main>', '<main><p><a href="../index.html">전체 보고서</a> · 전체 분석에서 선택한 계정 비교입니다.</p>', 1))
        private_write(output / filename, content)
        links.append((filename, '계정 비교 ' + str(index) + ' · ' + str(len(pairs)) + '쌍'))
    rows = ''.join('<li><a href="' + html.escape(path, quote=True) + '">' + html.escape(label) + '</a></li>' for path, label in links)
    private_write(output / 'index.html', '<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; base-uri \'none\'; form-action \'none\'"><title>Passive Scan · 분할 보고서</title><style>body{max-width:850px;margin:4em auto;padding:0 1em;font:17px/1.7 system-ui}li{margin:.5em 0}</style><h1>Passive Scan · 분할 보고서</h1><p>전체 ' + str(len(records)) + '개 기록을 페이지로 나누었습니다. 모든 증거와 전체 비교·흐름은 아래 파일에 보존됩니다. 분할 페이지는 검토 변경 JSON을 각각 내보낼 수 있습니다.</p><p><a href="requests.jsonl">전체 요청 JSONL</a> · <a href="analysis.json">전체 분석 JSON</a> · <a href="manifest.json">파일 영수증</a></p><ul>' + rows + '</ul></html>')


def _owned_report_files(output):
    manifest_path = output / 'manifest.json'
    if not manifest_path.is_file() or manifest_path.is_symlink():
        return set()
    try:
        manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
        files = manifest['files']
        if not isinstance(files, dict): return set()
        owned = set()
        for name, receipt in files.items():
            path = Path(name)
            if (not isinstance(name, str) or path.is_absolute() or '..' in path.parts or not re.fullmatch('[0-9a-f]{64}', str(receipt))):
                raise ValueError('Invalid report ownership manifest')
            if not (name in {'requests.jsonl','analysis.json','postman_review.json','review_queue.json','index.html','request_inventory.xlsx','capture_events.jsonl'} or re.fullmatch(r'evidence/[0-9a-f]{64}\.json',name) or re.fullmatch(r'pages/(?:requests|comparisons)-[0-9]+\.html',name)):
                raise ValueError('Unknown report-owned artifact')
            owned.add(name)
        if manifest.get('owner') == 'passive-scan-report':
            return owned | {'manifest.json'}
        # Adopt known v8.4 reports only after validating their source receipts.
        if manifest.get('schema_version') == '8.4' and set(files) == {'requests.jsonl', 'analysis.json'}:
            if any(not (output / name).is_file() or _file_digest(output / name) != receipt for name, receipt in files.items()):
                raise ValueError('Legacy report receipts do not match')
            owned.update(name for name in ('index.html', 'postman_review.json', 'review_queue.json', 'request_inventory.xlsx') if (output / name).is_file())
            evidence = output / 'evidence'
            if evidence.is_symlink(): raise ValueError('Report evidence must not be a symlink')
            if evidence.is_dir():
                for path in evidence.iterdir():
                    if path.is_file() and re.fullmatch(r'[0-9a-f]{64}\.json', path.name):
                        record = json.loads(path.read_text(encoding='utf-8'))
                        if record.get('observation_id') == path.stem:
                            owned.add(path.relative_to(output).as_posix())
            return owned | {'manifest.json'}
    except (KeyError, TypeError, json.JSONDecodeError) as error:
        raise ValueError('Invalid report ownership manifest') from error
    return set()


def _safe_report_path(output, name):
    target = output / name
    for path in (target, *target.parents):
        if path == output.parent: break
        if path.is_symlink(): raise ValueError('Report paths must not be symlinks')
    return target


def _carry_unowned(source, staging, owned, prefix=''):
    for item in source.iterdir():
        name=prefix+item.name
        if name in owned: continue
        target=staging/item.name
        if item.is_symlink():
            os.symlink(os.readlink(item),target)
        elif item.is_dir():
            target.mkdir(exist_ok=True,mode=0o700)
            _carry_unowned(item,target,owned,name+'/')
        elif item.is_file():
            try: os.link(item,target)
            except OSError: shutil.copy2(item,target)
        else:
            raise ValueError('Report contains an unsupported unrelated file; use a fresh output directory')


def _recover_report(output, journal):
    if not journal.exists(): return
    if journal.is_symlink(): raise ValueError('Invalid report recovery journal')
    value=json.loads(journal.read_text(encoding='utf-8'))
    if value.get('owner')!='passive-scan-report' or value.get('output')!=str(output):
        raise ValueError('Invalid report recovery journal')
    stage=Path(value['staging']); backup=Path(value['backup'])
    for path,label in ((stage,'stage'),(backup,'backup')):
        if path.parent!=output.parent or not path.name.startswith('.'+output.name+'-'+label+'-') or path.is_symlink():
            raise ValueError('Invalid report recovery path')
    if output.is_symlink(): raise ValueError('Report output must not be a symlink')
    if output.exists():
        current=None
        try: current=json.loads((output/'manifest.json').read_text(encoding='utf-8')).get('run_id')
        except (OSError,ValueError): pass
        if current==value.get('run_id'):
            if backup.exists(): shutil.rmtree(backup)
        elif backup.exists():
            raise ValueError('Report recovery conflict; previous output remains in backup')
        # Otherwise the old directory was never moved, or was already restored.
    elif backup.exists():
        os.replace(backup,output)
    else:
        raise ValueError('Report recovery requires the saved output directory')
    if stage.exists(): shutil.rmtree(stage)
    journal.unlink()


def _acquire_report_lock(fd):
    from .locking import acquire_file_lock
    acquire_file_lock(fd,'Report generation is already in progress')


def generate(records, output, xlsx=False, ignore_fields=(), capture_events=None, resource_headers=()):
    """Publish complete generations under a writer lock, with crash recovery.

    Directory replacement takes two renames on portable filesystems. The journal
    restores the old generation after an interrupted swap; it is not an atomic
    directory exchange. No mixed generation is published and backups survive
    failed recovery. Unrelated files are carried into the prepared directory.
    """
    output=Path(output).absolute()
    if output.is_symlink() or (output.exists() and not output.is_dir()):
        raise ValueError('Report output must be a directory, not a symlink')
    output.parent.mkdir(parents=True,exist_ok=True,mode=0o700)
    lock_path=output.parent/('.'+output.name+'.report.lock')
    journal=output.parent/('.'+output.name+'.report-journal.json')
    if lock_path.is_symlink(): raise ValueError('Report lock must not be a symlink')
    fd=os.open(lock_path,os.O_CREAT|os.O_RDWR|getattr(os,'O_NOFOLLOW',0),0o600)
    staging=None
    try:
        _acquire_report_lock(fd)
        _recover_report(output,journal)
        if capture_events is not None: capture_events=list(capture_events)
        staging=Path(tempfile.mkdtemp(prefix='.'+output.name+'-stage-',dir=output.parent))
        count=_generate(records,staging,xlsx,ignore_fields,capture_events,resource_headers)
        if not output.exists():
            os.replace(staging,output)
            return count
        owned=_owned_report_files(output)
        new_files={path.relative_to(staging).as_posix() for path in staging.rglob('*') if path.is_file()}
        for name in owned|new_files:
            target=_safe_report_path(output,name)
            if target.exists() and (not target.is_file() or name not in owned):
                raise ValueError('Report would overwrite an unowned file; use a fresh output directory')
        _carry_unowned(output,staging,owned)
        backup=Path(tempfile.mkdtemp(prefix='.'+output.name+'-backup-',dir=output.parent))
        backup.rmdir()
        run_id=json.loads((staging/'manifest.json').read_text(encoding='utf-8'))['run_id']
        private_write(journal,json.dumps({'owner':'passive-scan-report','output':str(output),'staging':str(staging),'backup':str(backup),'run_id':run_id}))
        try:
            os.replace(output,backup)
            os.replace(staging,output)
        except BaseException:
            # A failed restore keeps both the journal and backup for the next run.
            if not output.exists() and backup.exists(): os.replace(backup,output)
            if staging.exists(): shutil.rmtree(staging)
            journal.unlink()
            raise
        _recover_report(output,journal)
        return count
    finally:
        if staging is not None and staging.exists() and not journal.exists(): shutil.rmtree(staging)
        os.close(fd)
